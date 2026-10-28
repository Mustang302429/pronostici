"""Genera il report delle prossime partite: py -m motore.genera"""
import warnings; warnings.filterwarnings("ignore")
import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule
from .config import OUT, CAMPIONATI, COPPE
from . import scarica, presente, archivio
from .dati import carica
from .modello import Modello
from .calendario import prossime
from .contesto import scheda
from .backtest import metriche, calibrazione

NOMI_COMP = {**CAMPIONATI, **COPPE}
CANDIDATI = ["1", "X", "2", "1X", "X2", "Over 2.5", "Under 2.5", "Over 3.5", "Gol", "NoGol"]

def costruisci(oggi=None, giorni=11, aggiorna=True, ramo2=True):
    if aggiorna:
        scarica.RAW.mkdir(parents=True, exist_ok=True)
        print("[1/5] Scarico i campionati (football-data.co.uk)...", flush=True)
        scarica.campionati().to_csv(scarica.RAW / "campionati.csv", index=False)
        print("[2/5] Scarico le coppe europee (ESPN)...", flush=True)
        scarica.coppe().to_csv(scarica.RAW / "coppe.csv", index=False)
        scarica.calendario_campionati().to_csv(scarica.RAW / "calendario.csv", index=False)
    oggi = pd.Timestamp(oggi or pd.Timestamp.now(tz="Europe/Rome").date())
    print("[3/5] Ramo 1: calcolo la forza delle squadre...", flush=True)
    df, _ = carica()
    m = Modello(emivita=400, ridge=2).fit(df, oggi + pd.Timedelta(days=1))
    cal = prossime(giorni, oggi)
    if ramo2:
        print(f"[4/5] Ramo 2: leggo infortuni, squalifiche e formazioni di {len(cal)} partite (FotMob)...", flush=True)
        cal = presente.analizza(cal)
    else:
        cal["corr_casa"] = cal["corr_ospite"] = 1.0
        cal["assenti_casa"] = cal["assenti_ospite"] = ""; cal["formazione"] = "ramo 2 disattivato"
    print("[5/5] Calcolo i pronostici e scrivo il report...", flush=True)
    righe, schede, arch = [], [], []
    for r in cal.itertuples():
        p1 = m.mercati(r.casa, r.ospite, r.comp, r.neutro)
        p2 = m.mercati(r.casa, r.ospite, r.comp, r.neutro, r.corr_casa, r.corr_ospite)
        riga = {"Data": r.data, "Competizione": NOMI_COMP.get(r.comp, r.comp),
                "Casa": r.casa_espn, "Ospite": r.ospite_espn,
                "xG casa": p2["xG_casa"], "xG ospite": p2["xG_ospite"]}
        for k in ["1", "X", "2", "1X", "X2", "Over 1.5", "Over 2.5", "Under 2.5", "Over 3.5", "Gol", "NoGol"]:
            riga[k] = p2[k]
        segno = max("1X2", key=lambda k: p2[k])
        riga["Segno favorito"] = f"{segno} ({p2[segno]:.0%})"
        best = max(CANDIDATI, key=lambda k: p2[k])
        riga["Esito più probabile"] = f"{best} ({p2[best]:.0%})"
        riga["Risultati esatti probabili"] = p2["Risultati probabili"]
        # effetto del presente sul segno 1X2
        d1 = p2["1"] - p1["1"]; d2 = p2["2"] - p1["2"]
        riga["Effetto assenze"] = (f"1 {d1:+.0%} / 2 {d2:+.0%}" if max(abs(d1), abs(d2)) >= 0.01 else "trascurabile")
        riga["Solo statistica 1-X-2"] = f"{p1['1']:.0%}-{p1['X']:.0%}-{p1['2']:.0%}"
        q = [getattr(r, x, np.nan) for x in ("q1", "qx", "q2")]
        riga["Bookmaker 1"] = riga["Bookmaker X"] = riga["Bookmaker 2"] = None
        riga["Disaccordo col mercato"] = ""
        if all(pd.notna(q)):
            imp = 1 / np.array(q, dtype=float); imp /= imp.sum()
            riga["Bookmaker 1"], riga["Bookmaker X"], riga["Bookmaker 2"] = imp
            diff = np.array([p2["1"], p2["X"], p2["2"]]) - imp
            i = int(np.argmax(np.abs(diff)))
            if abs(diff[i]) >= 0.08: riga["Disaccordo col mercato"] = f"{'1X2'[i]}: {diff[i]:+.0%}"
        riga["Formazioni"] = r.formazione
        riga["Assenti casa"] = r.assenti_casa; riga["Assenti ospite"] = r.assenti_ospite
        riga["Affidabilità dati"] = p1["affidabilita"]
        righe.append(riga)
        schede.append({"Data": r.data, "Partita": f"{r.casa_espn} - {r.ospite_espn}",
                       "Competizione": NOMI_COMP.get(r.comp, r.comp),
                       **scheda(df, r.casa, r.ospite, oggi + pd.Timedelta(days=1))})
        arch.append({"data": r.data, "comp": r.comp, "casa": r.casa, "ospite": r.ospite,
                     **{f"R1_{k}": p1[v] for k, v in [("1", "1"), ("X", "X"), ("2", "2"), ("O25", "Over 2.5"), ("GG", "Gol")]},
                     **{f"R2_{k}": p2[v] for k, v in [("1", "1"), ("X", "X"), ("2", "2"), ("O25", "Over 2.5"), ("GG", "Gol")]},
                     "corr_casa": r.corr_casa, "corr_ospite": r.corr_ospite, "q1": q[0], "qx": q[1], "q2": q[2]})
    archivio.salva(pd.DataFrame(arch))
    _, riep = archivio.verifica(df)
    return pd.DataFrame(righe), pd.DataFrame(schede), riep

# ------------------------------------------------------------------ Excel
H = PatternFill("solid", fgColor="1F3A5F"); W = Font(color="FFFFFF", bold=True); THIN = Side(style="thin", color="D0D7E2")
PERC = ["1", "X", "2", "1X", "X2", "Over 1.5", "Over 2.5", "Under 2.5", "Over 3.5", "Gol", "NoGol",
        "Bookmaker 1", "Bookmaker X", "Bookmaker 2"]

def _foglio(ws, d, perc=(), larg=None, freeze="C2"):
    ws.append(list(d.columns))
    for c in ws[1]: c.fill, c.font, c.alignment = H, W, Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in d.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and np.isnan(v)) else v for v in row])
    for j, col in enumerate(d.columns, 1):
        L = get_column_letter(j)
        ws.column_dimensions[L].width = (larg or {}).get(col, max(10, min(45, len(str(col)) + 4)))
        for c in ws[L][1:]:
            c.border = Border(bottom=THIN)
            if col in perc and c.value is not None: c.number_format = "0%"
            elif isinstance(c.value, float): c.number_format = "0.00"
            elif isinstance(c.value, pd.Timestamp): c.number_format = "dd/mm hh:mm"
    ws.freeze_panes = freeze; ws.auto_filter.ref = ws.dimensions; ws.row_dimensions[1].height = 32

def salva_excel(prev, sch, bt=None, riep=None, nome=None):
    OUT.mkdir(exist_ok=True, parents=True)
    nome = nome or OUT / f"pronostici_{pd.Timestamp.now():%Y-%m-%d}.xlsx"
    wb = Workbook()
    ws = wb.active; ws.title = "Prossime partite"
    _foglio(ws, prev, PERC, {"Data": 13, "Competizione": 22, "Casa": 22, "Ospite": 22, "Segno favorito": 14,
                             "Esito più probabile": 18, "Risultati esatti probabili": 30, "Effetto assenze": 18,
                             "Solo statistica 1-X-2": 16, "Disaccordo col mercato": 14, "Formazioni": 16,
                             "Assenti casa": 60, "Assenti ospite": 60}, freeze="E2")
    for c in PERC[:11]:
        L = get_column_letter(list(prev.columns).index(c) + 1)
        ws.conditional_formatting.add(f"{L}2:{L}{len(prev)+1}", ColorScaleRule(start_type="num", start_value=0.2,
            start_color="FFFFFF", end_type="num", end_value=0.85, end_color="5DB37E"))
    _foglio(wb.create_sheet("Schede partita"), sch, larg={"Partita": 36, "Precedenti (ult.5)": 90, "Competizione": 22})
    ws5 = wb.create_sheet("Verifica pronostici")
    ws5.append(["Verifica automatica: pronostici emessi da questo sistema confrontati con i risultati reali"])
    ws5["A1"].font = Font(bold=True, size=13)
    if riep is not None and len(riep):
        ws5.append(list(riep.columns))
        for c in ws5[2]: c.fill, c.font = H, W
        for row in riep.itertuples(index=False): ws5.append(list(row))
        for r in range(3, 3 + len(riep)):
            for cc in (3, 4, 5): ws5.cell(r, cc).number_format = "0.0%"
            ws5.cell(r, 6).number_format = "0.000"
    else:
        ws5.append(["Ancora nessuna partita verificata: i numeri compariranno dopo le prime partite giocate."])
    for L, w in zip("ABCDEF", (30, 18, 22, 24, 20, 30)): ws5.column_dimensions[L].width = w
    if bt is not None:
        ws3 = wb.create_sheet("Affidabilità (backtest)")
        met = metriche(bt); cal = calibrazione(bt).reset_index()
        ws3.append(["Backtest stagione 2025/26 (Ramo 1): ogni partita prevista usando SOLO i dati disponibili prima"]); ws3["A1"].font = Font(bold=True, size=13)
        for a, b in [("Partite testate", met["partite"]), ("Il favorito del modello ha vinto", met["accuratezza_favorito"]),
                     ("Over/Under 2.5 indovinati", met["accuratezza_over25"]), ("Gol/NoGol indovinati", met["accuratezza_GG/NG"]),
                     ("Log-loss 1X2 modello (più basso = meglio)", met["logloss_1X2_modello_(stesse partite)"]),
                     ("Log-loss 1X2 quote bookmaker in chiusura", met["logloss_1X2_bookmaker"])]:
            ws3.append([a, b])
        for r in range(3, 6): ws3.cell(r, 2).number_format = "0.0%"
        ws3.append([]); ws3.append(["Calibrazione: quando il modello dice X%, quanto spesso succede davvero?"]); ws3.cell(ws3.max_row, 1).font = Font(bold=True)
        ws3.append(["Fascia di probabilità", "Previsto", "Successo reale", "N. casi"])
        for c in ws3[ws3.max_row]: c.fill, c.font = H, W
        for r in cal.itertuples(index=False):
            ws3.append([str(r.fascia), r.previsto, r.reale, int(r.n)])
            ws3.cell(ws3.max_row, 2).number_format = ws3.cell(ws3.max_row, 3).number_format = "0.0%"
        ws3.column_dimensions["A"].width = 48; ws3.column_dimensions["B"].width = 14; ws3.column_dimensions["C"].width = 16
    ws4 = wb.create_sheet("Come leggere", 0)
    for t in [
        "COME LEGGERE IL REPORT",
        "",
        "RAMO 1 - STATISTICA: storico dal 2022 (football-data.co.uk + ESPN), modello Poisson/Dixon-Coles.",
        "RAMO 2 - PRESENTE: infortuni, squalifiche e formazioni probabili (FotMob). Ogni giocatore pesa in base al",
        "   valore di mercato e al voto stagionale. Se mancano titolari importanti in attacco, i gol attesi della squadra",
        "   scendono; se mancano in difesa, salgono quelli dell'avversario (correzione massima ±25%).",
        "",
        "Le probabilità nelle colonne 1/X/2/Over/Gol... sono già CORRETTE dal ramo 2 (verde = più probabile).",
        "Solo statistica 1-X-2: le probabilità del solo ramo 1, per confronto.",
        "Effetto assenze: di quanto il ramo 2 ha spostato le probabilità di 1 e di 2.",
        "Formazioni: 'probabile'/'ufficiale' = usata la formazione annunciata; 'da indisponibili' = formazione non",
        "   ancora uscita, si toglie dalla rosa chi è infortunato o squalificato (in dubbio = conta a metà).",
        "Assenti: i 5 assenti più pesanti per squadra, con reparto (POR/DIF/CEN/ATT) e gol stagionali.",
        "Bookmaker 1/X/2: probabilità implicite nelle quote medie (margine tolto). Disaccordo col mercato: differenza",
        "   di almeno 8 punti tra sistema e bookmaker -> conviene controllare le notizie.",
        "Affidabilità dati: ALTA = entrambe le squadre hanno storico nei campionati coperti.",
        "Verifica pronostici: ogni pronostico viene archiviato e confrontato col risultato reale (Ramo 1 vs Ramo 2).",
        "",
        "Consiglio: rilancia il report il giorno della partita, quando escono le formazioni probabili.",
        "Limite onesto: il sistema stima probabilità, non certezze; da solo non batte le quote di chiusura.",
    ]:
        ws4.append([t])
    ws4["A1"].font = Font(bold=True, size=13); ws4.column_dimensions["A"].width = 120
    wb.save(nome); return nome

if __name__ == "__main__":
    import sys
    from . import sito
    prev, sch, riep = costruisci()
    sito.scrivi(prev, sch, riep, sito.riepilogo_backtest())
    if "--solo-sito" in sys.argv:
        print(f"\nFATTO. {len(prev)} partite. Dati web aggiornati in docs/data.json"); raise SystemExit
    try:
        bt = pd.read_csv(OUT / "backtest_2025_26.csv")
    except FileNotFoundError:
        bt = None
    f = salva_excel(prev, sch, bt, riep)
    print(f"\nFATTO. {len(prev)} partite. File creato:\n{f}")
